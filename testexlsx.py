from openpyxl import load_workbook
import warnings
warnings.filterwarnings("ignore")

wbFonte = load_workbook('F:\\Work\\Ementas\\ementa7_20_07_2016.xlsx')
wsFonte = wbFonte.worksheets[0]

wbTarget = load_workbook('F:\\Work\\Ementas\\teste.xlsx')
wsTarget = wbTarget.worksheets[0]

#FONTE
s1BerF = 	[5,6,8,9,10]
s2BerF = 	[12,13,14,15,16,17]
sParqueF = 	[19,20,21,22,23,24]
sPreCatlF = [26,27,28,29,30,31]
dietaF = 	[33,34,35,36]
letrasF = 	['B','E','H','K','N']


#TARGET
s1BerT = 	[6,7,9,10,11]
s2BerT = 	[15,16,17,18,19,20]
sParqueT = 	[24,25,26,27,28,29]
sPreCatlT =     [33,34,35,36,37,38]
dietaT = 	[42,43,44,45]
letrasT = 	['B','C','D','E','F']
#salas = 	[s1BerT,s2Ber,sParque,sPreCatl,dieta]

#LOGOTIPO
logo = "F:\\Raul\Work\Ementas\Logotipo_alta_png.png"

#SALA BER1
for i in range(len(letrasF)):
	for k in range(len(s1BerF)):
			wsTarget["%s%d" % (letrasT[i],s1BerT[k])].value = wsFonte["%s%d" % (letrasF[i],s1BerF[k])].value
				
#SALA BER2
for i in range(len(letrasF)):
	for k in range(len(s2BerF)):
			wsTarget["%s%d" % (letrasT[i],s2BerT[k])].value = wsFonte["%s%d" % (letrasF[i],s2BerF[k])].value

#SALA PARQUE
for i in range(len(letrasF)):
	for k in range(len(sParqueF)):
			wsTarget["%s%d" % (letrasT[i],sParqueT[k])].value = wsFonte["%s%d" % (letrasF[i],sParqueF[k])].value

#PRE/CATL
for i in range(len(letrasF)):
	for k in range(len(sPreCatlF)):
			wsTarget["%s%d" % (letrasT[i],sPreCatlT[k])].value = wsFonte["%s%d" % (letrasF[i],sPreCatlF[k])].value
#DIETA
for i in range(len(letrasF)):
	for k in range(len(dietaF)):
			wsTarget["%s%d" % (letrasT[i],dietaT[k])].value = wsFonte["%s%d" % (letrasF[i],dietaF[k])].value
		
wbTarget.save('teste7_20_07_2016.xlsx')
